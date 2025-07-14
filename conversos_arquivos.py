# converter_app.py
# pip install pandas openpyxl xlwt python-docx docx2pdf pdf2docx customtkinter

import customtkinter as ctk
import threading
import pandas as pd
import xlwt
from docx import Document
from docx2pdf import convert as docx2pdf_convert
from pdf2docx import Converter as PDF2DOCXConverter
from tkinter import filedialog, messagebox

# Configura tema moderno
ctk.set_appearance_mode("System")  # System, Light, Dark
ctk.set_default_color_theme("green")  # blue (default), dark-blue, green

class SplashScreen:
    def __init__(self, parent):
        self.root = ctk.CTkToplevel(parent)
        self.root.overrideredirect(True)
        # Centraliza
        w, h = 450, 250
        sw = parent.winfo_screenwidth(); sh = parent.winfo_screenheight()
        x = (sw - w) // 2; y = (sh - h) // 2
        self.root.geometry(f"{w}x{h}+{x}+{y}")
        logo = ctk.CTkLabel(self.root, text="Conversor de arquivos", font=ctk.CTkFont(size=30, weight="bold"))
        logo.pack(expand=True)
        self.spinner = ctk.CTkLabel(self.root, text="", font=ctk.CTkFont(size=20))
        self.spinner.pack(pady=10)
        self.chars = ["⟳","⟲","⟳","⟲"]
        self.idx = 0; self.running = True
        self.animate(); self.root.after(2500, self.close)

    def animate(self):
        if not self.running: return
        self.spinner.configure(text=self.chars[self.idx % len(self.chars)])
        self.idx += 1
        self.root.after(100, self.animate)

    def close(self):
        self.running = False; self.root.destroy()

class ConverterApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Bang Top Converter")
        self.geometry("600x400")
        self.grid_rowconfigure(0, weight=1); self.grid_columnconfigure(0, weight=1)

        # Frame principal sem 'padding' no construtor
        frame = ctk.CTkFrame(self, corner_radius=15)
        frame.grid(row=0, column=0, padx=20, pady=20, sticky="nsew")
        frame.grid_columnconfigure((0,1), weight=1)

        ctk.CTkLabel(frame, text="Escolha a conversão:", font=ctk.CTkFont(size=16)).grid(row=0, column=0, columnspan=2, sticky="w")
        self.combo = ctk.CTkComboBox(frame, values=[
            "CSV → XLSX","CSV → DOCX","DOCX → PDF","XLSX → XLS","XLS → XLSX","PDF → DOCX"
        ])
        self.combo.set("CSV → XLSX")
        self.combo.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(10,10))

        self.input_entry = ctk.CTkEntry(frame, placeholder_text="Arquivo de entrada")
        self.input_entry.grid(row=2, column=0, sticky="ew", padx=(0,10))
        ctk.CTkButton(frame, text="Selecionar", command=self.select_input).grid(row=2, column=1, sticky="ew")

        self.output_entry = ctk.CTkEntry(frame, placeholder_text="Arquivo de saída")
        self.output_entry.grid(row=3, column=0, sticky="ew", padx=(0,10), pady=(10,0))
        ctk.CTkButton(frame, text="Selecionar", command=self.select_output).grid(row=3, column=1, sticky="ew", pady=(10,0))

        self.convert_btn = ctk.CTkButton(frame, text="Converter", command=self.start)
        self.convert_btn.grid(row=4, column=0, columnspan=2, pady=(20,0), sticky="ew")

        self.loading_label = ctk.CTkLabel(frame, text="", font=ctk.CTkFont(size=24))
        self.loading_label.grid(row=5, column=0, columnspan=2, pady=(20,0))

        self.input_path = ""; self.output_path = ""
        self.loading_chars = ["⏳","🔄","💫","🔃"]; self.load_idx = 0

    def select_input(self):
        path = filedialog.askopenfilename(title="Selecione entrada")
        if path:
            self.input_path = path
            self.input_entry.delete(0, "end")
            self.input_entry.insert(0, path)

    def select_output(self):
        ext = self.get_ext()
        path = filedialog.asksaveasfilename(defaultextension=ext, filetypes=[("Arquivo", f"*{ext}")], title="Selecione destino")
        if path:
            self.output_path = path
            self.output_entry.delete(0, "end")
            self.output_entry.insert(0, path)

    def get_ext(self):
        return {
            "CSV → XLSX": ".xlsx","CSV → DOCX": ".docx","DOCX → PDF": ".pdf",
            "XLSX → XLS": ".xls","XLS → XLSX": ".xlsx","PDF → DOCX": ".docx"
        }[self.combo.get()]

    def start(self):
        if not self.input_path or not self.output_path:
            messagebox.showwarning("Oops","Informe input e output!")
            return
        self.convert_btn.configure(state=ctk.DISABLED)
        threading.Thread(target=self.run_conversion).start()
        self.animate_loading()

    def animate_loading(self):
        if self.convert_btn._state == "normal":
            self.loading_label.configure(text=""); return
        self.loading_label.configure(text=self.loading_chars[self.load_idx % len(self.loading_chars)])
        self.load_idx += 1; self.after(100, self.animate_loading)

    def run_conversion(self):
        try:
            funcs = {
                "CSV → XLSX": self.csv_to_xlsx,
                "CSV → DOCX": self.csv_to_docx,
                "DOCX → PDF": self.docx_to_pdf,
                "XLSX → XLS": self.xlsx_to_xls,
                "XLS → XLSX": self.xls_to_xlsx,
                "PDF → DOCX": self.pdf_to_docx
            }
            funcs[self.combo.get()]()
            messagebox.showinfo("Sucesso","Conversão concluída!")
        except Exception as e:
            messagebox.showerror("Erro", str(e))
        finally:
            self.convert_btn.configure(state=ctk.NORMAL)

    def csv_to_xlsx(self):
        df = pd.read_csv(self.input_path); df.to_excel(self.output_path, index=False)

    def csv_to_docx(self):
        df = pd.read_csv(self.input_path)
        doc = Document()
        table = doc.add_table(rows=1, cols=len(df.columns))
        for i, col in enumerate(df.columns): table.cell(0,i).text = col
        for row in df.itertuples(index=False):
            cells = table.add_row().cells
            for i, val in enumerate(row): cells[i].text = str(val)
        doc.save(self.output_path)

    def docx_to_pdf(self):
        docx2pdf_convert(self.input_path, self.output_path)

    def xlsx_to_xls(self):
        df = pd.read_excel(self.input_path)
        wb = xlwt.Workbook(); ws = wb.add_sheet('Sheet1')
        for r, row in enumerate(df.values):
            for c, val in enumerate(row): ws.write(r, c, val)
        wb.save(self.output_path)

    def xls_to_xlsx(self):
        df = pd.read_excel(self.input_path); df.to_excel(self.output_path, index=False)

    def pdf_to_docx(self):
        cv = PDF2DOCXConverter(self.input_path); cv.convert(self.output_path); cv.close()

if __name__ == "__main__":
    app = ConverterApp()
    SplashScreen(app)
    app.after(2600, lambda: app.deiconify())
    app.withdraw()
    app.mainloop()
